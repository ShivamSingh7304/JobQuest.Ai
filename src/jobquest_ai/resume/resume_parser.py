
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 10


class ResumeParsingError(Exception):
    """Raised when a resume cannot be parsed."""


def validate_resume(file_path: str | Path) -> Path:
    """Validate the resume path, extension, and file size."""

    path = Path(file_path)

    if not path.exists():
        raise ResumeParsingError("Resume file does not exist.")

    if not path.is_file():
        raise ResumeParsingError("The provided path is not a file.")

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ResumeParsingError(
            "Unsupported file format. Upload a PDF or DOCX file."
        )

    if path.stat().st_size == 0:
        raise ResumeParsingError("The uploaded resume is empty.")

    max_size = MAX_FILE_SIZE_MB * 1024 * 1024

    if path.stat().st_size > max_size:
        raise ResumeParsingError(
            f"Resume exceeds the {MAX_FILE_SIZE_MB} MB limit."
        )

    return path


def extract_pdf_text(file_path: Path) -> str:
    """Extract text from a PDF resume."""

    try:
        reader = PdfReader(str(file_path))

        if reader.is_encrypted:
            raise ResumeParsingError(
                "Password-protected PDFs are not supported."
            )

        pages = []

        for page in reader.pages:
            pages.append(page.extract_text() or "")

        return "\n".join(pages).strip()

    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read the PDF resume."
        ) from exc


def extract_docx_text(file_path: Path) -> str:
    """Extract paragraphs and table content from a DOCX resume."""

    try:
        document = Document(str(file_path))
        content = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        # Resume templates often put skills or experience in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    content.append(" | ".join(cells))

        return "\n".join(content).strip()

    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read the DOCX resume."
        ) from exc


def parse_resume(file_path: str | Path) -> dict:
    """
    Validate a resume, extract its text, and return structured metadata.

    Returns:
        {
            "file_name": str,
            "file_type": str,
            "text": str,
            "character_count": int
        }
    """

    path = validate_resume(file_path)

    if path.suffix.lower() == ".pdf":
        text = extract_pdf_text(path)
    else:
        text = extract_docx_text(path)

    if not text:
        raise ResumeParsingError(
            "No readable text found. The PDF may be scanned or image-based."
        )

    return {
        "file_name": path.name,
        "file_type": path.suffix.lower().lstrip("."),
        "text": text,
        "character_count": len(text),
    }
