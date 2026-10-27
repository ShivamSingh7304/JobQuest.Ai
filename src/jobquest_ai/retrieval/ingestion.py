from pathlib import Path
from typing import Any
from docx import Document as DocxDocument
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentIngestionError(Exception):
    """Raised when document ingestion fails."""


def _extract_file_documents(
    file_path: Path,
    extra_metadata: dict[str, Any] | None = None,
) -> list[Document]:
    """Extract text from one supported file."""

    suffix = file_path.suffix.lower()
    metadata = {
        "source": str(file_path.resolve()),
        "file_name": file_path.name,
    }
    metadata.update(extra_metadata or {})

    documents = []

    try:
        if suffix == ".pdf":
            reader = PdfReader(str(file_path))

            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""

                if text.strip():
                    page_metadata = {
                        **metadata,
                        "page": page_number,
                    }
                    documents.append(
                        Document(
                            page_content=text,
                            metadata=page_metadata,
                        )
                    )

        elif suffix == ".docx":
            doc = DocxDocument(str(file_path))
            text = "\n".join(
                paragraph.text
                for paragraph in doc.paragraphs
                if paragraph.text.strip()
            )

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata=metadata.copy(),
                    )
                )

        elif suffix in {".txt", ".md"}:
            text = file_path.read_text(
                encoding="utf-8",
                errors="replace",
            )

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata=metadata.copy(),
                    )
                )

        else:
            raise DocumentIngestionError(
                f"Unsupported file type: {suffix}"
            )

    except DocumentIngestionError:
        raise
    except Exception as exc:
        raise DocumentIngestionError(
            f"Could not read {file_path.name}: {exc}"
        ) from exc

    return documents


def load_documents(
    source_path: str | Path,
    metadata: dict[str, Any] | None = None,
) -> list[Document]:
    """
    Load a supported file or all supported files in a directory.

    Optional metadata can include fields such as technology or version.
    """
    path = Path(source_path)

    if not path.exists():
        raise DocumentIngestionError(
            f"Source path does not exist: {path}"
        )

    if path.is_file():
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise DocumentIngestionError(
                f"Unsupported file type: {path.suffix}"
            )
        files = [path]

    elif path.is_dir():
        files = sorted(
            file
            for file in path.rglob("*")
            if file.is_file()
            and file.suffix.lower() in SUPPORTED_EXTENSIONS
        )

    else:
        raise DocumentIngestionError(
            f"Invalid source path: {path}"
        )

    documents = []

    for file_path in files:
        documents.extend(
            _extract_file_documents(file_path, metadata)
        )

    if not documents:
        raise DocumentIngestionError(
            "No extractable text was found in the supplied source."
        )

    return documents


def split_documents(
    documents: list[Document],
    chunk_size: int = 1000,
    chunk_overlap: int = 150,
) -> list[Document]:
    """Split loaded documents into overlapping chunks."""

    if not documents:
        return []

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError(
            "chunk_overlap must be non-negative and smaller "
            "than chunk_size."
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        add_start_index=True,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    return splitter.split_documents(documents)


def load_and_split_documents(
    source_path: str | Path,
    chunk_size: int = 1000,
    chunk_overlap: int = 150,
    metadata: dict[str, Any] | None = None,
) -> list[Document]:
    """Convenience function for the complete ingestion step."""

    documents = load_documents(
        source_path=source_path,
        metadata=metadata,
    )

    return split_documents(
        documents,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
